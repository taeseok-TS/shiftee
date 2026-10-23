# 큐브티 전자계약 양식 2건(2026-10-23, 본부 요청) — 모두싸인 양식 PDF 를 큐브티 워드 양식으로 옮긴다.
#
#  1) 「주말 학습실장 근로계약서」 — 원본: 모두싸인 「주말 학습실장 근로계약서_지점_이름」(PDF 생성일 2026-09-01 20:46, 9/1 판).
#     본부 요구 「9/1 판, 포괄임금 문구 없이」 — 이 원본에는 평일 학습실장(계약직, 3/4 판) 제4조 라.2
#     「월 급여 합계액은 한 달을 4.2주로 가산하여 산정한 포괄임금이며, …고정 지급한다.」가 처음부터 없다(옮길 것도 뺄 것도 없음).
#     문구는 원본 그대로, 빈칸만 {칸} 으로 바꾼다. 쪽 나눔(3쪽)도 원본과 같게.
#     ※ 원본 PDF 글자 추출의 「제 1 조」「1 일」은 진짜 빈칸이 아니라 워드의 한글·숫자 자동 간격(2.5pt, 빈칸은 3.5pt)이다 → 「제 1조」「1일」로 쓴다.
#     10/23 본부(이예지) 답 반영 — 문장 내용은 그대로, 띄어쓰기·형식만 바로잡음:
#       「근로 계약을」→「근로계약을」, 「"갑" 의」→「"갑"의」, 「부적합 하다고」→「부적합하다고」, 「않 을」→「않을」,
#       제9조 가.·나. 한 문단 → 두 문단, 제3조 다. 들여쓰기를 가.·나.와 같게, 제10조 「갑의 제규정」→「"갑"의 제규정」
#       요일 표 끝에 일요일 칸 추가, ④ 월 급여합계 = 서버 파생 칸 {주말급여합계}(①+②+③). 원장 줄은 「대치 재수종합」 고정.
#       결재 순서는 본부 → 원장 → 근로자(직인 {대표서명}·《원장서명》 자리는 문서에 그대로).
#  2) 「코디 근로계약서(계약직 일반)」 오전·오후 두 판 — 큐브티 v1(7/31, 코디근로계약서_계약직일반_템플릿.docx)을 그대로 읽어
#     제4조 나. 한 문장만 바꾼다. 모두싸인 원본 두 판(일반 2025-05-08 / 오전시간 2026-07-20)의 차이는 이 문장 하나뿐이다(글자·픽셀 대조).
#       · 일반(오후) : 근무시간은 통상적인 근무일 기준 오후 ___ ～ 오후 ___ 로 한다.
#       · 오전시간   : 근무시간은 통상적인 근무일 기준 ___ ～ ___ 로 한다.
#     오후판은 「오후 {출근시각12} ～ 오후 {퇴근시각12}」(12시간제 파생 칸 — 「오후 2:00」), 오전판은 v1 그대로 {출근시각}/{퇴근시각}.
#     등록: 기존 7/31 양식 파일을 오전판으로 교체 + 오후판 하나 추가(목록 2개).
#     v1 은 제4조 다. 「"을"의 휴게시간은…」 한 줄만 들여쓰기가 빠져 있어(원본 두 판은 가.·나.와 같은 들여쓰기) 두 판 모두 맞춘다.
#
# 자리표시 이름은 기존 양식과 같게: {직원명} {지점} {원장명} {계약시작일} {계약종료일} {작성일} {월근로시간} {직무}
#   서명 마커 {대표서명}(직인) {근로자서명} 《원장서명》(build_tpl8 와 같은 글자 마커)
# 실행: python build_tpl_weekend_codi.py <출력 폴더> [<코디 v1 docx 경로>]
import sys, os
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else "."
CODI_V1 = sys.argv[2] if len(sys.argv) > 2 else os.path.expanduser(r"~\Downloads\코디근로계약서_계약직일반_템플릿.docx")
LOGO = os.path.join(HERE, "nexcube_logo.png")   # 큐브티 v1 코디 양식 바닥글의 NEXCUBE 로고(원본 PDF 들의 오른쪽 아래 로고)
FONT = "맑은 고딕"
C, R, L = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT

# 주말 학습실장 원장 줄 — 원본이 「대치 재수종합」으로 박혀 있다(파일 이름은 「지점_이름」). 다른 지점에도 쓰려면 "{지점명}" 으로
WEEKEND_BRANCH_TEXT = "대치 재수종합"
# 코디 오후판의 출근·퇴근 칸 — 서버 파생 칸(12시간제 숫자, 「14:00」→「2:00」, lib/contract-fields.ts). 작성 화면은 원천 칸
# {출근시각}/{퇴근시각}(시간 선택, 24시간 값)을 그대로 띄운다(fields API 가 원천 칸을 덧붙임). 오전판은 v1 그대로 {출근시각}/{퇴근시각}
PM_FIELDS = ("출근시각12", "퇴근시각12")

BR = "|BR|"      # 문단 안 줄바꿈 표시
LINE = 22.8      # 원본 줄 간격(pt) — 주말 원본 PDF 에서 잰 값
GAP = 17.3       # 조항 사이 빈 줄 몫(pt)
IND = 20.8       # 가./나. 들여쓰기(pt)
HANG = 15.7      # 「가. 」 너비 — 여러 줄 문단의 둘째 줄을 글 첫머리에 맞춘다


def set_font(r, size=None, bold=False, underline=False):
    r.bold = bold
    r.underline = underline
    if size: r.font.size = Pt(size)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return r


def para(d, pieces=(), align=None, indent=None, first=None, before=0, size=None, bold=False, line=LINE, container=None):
    """pieces: 문자열 또는 (문자열, 밑줄) 목록. 원본 밑줄 자리(빈칸·「매월 1 일부터 말일까지」 등)를 그대로 살린다"""
    p = (container or d).add_paragraph()
    if isinstance(pieces, str): pieces = [pieces]
    for pc in pieces:
        text, ul = (pc, False) if isinstance(pc, str) else pc
        for k, part in enumerate(text.split(BR)):   # BR = 원본 줄바꿈 자리(원본은 줄마다 끊긴 문단이라 같은 자리에서 끊는다)
            r = set_font(p.add_run(), size=size, bold=bold, underline=ul)
            if k: r.add_break()
            r.add_text(part)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(0)
    if line: pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(line)
    pf.widow_control = False   # 원본은 제9조 가.가 2쪽 끝 한 줄·3쪽 첫 줄로 갈린다 — 외톨이 줄 막기를 끄면 같은 자리에서 갈린다
    if align is not None: p.alignment = align
    if indent is not None: pf.left_indent = Pt(indent)
    if first is not None: pf.first_line_indent = Pt(first)
    return p


def head(d, text, before=GAP):
    return para(d, text, bold=True, before=before)


def borders(t, sz=8):
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        b.append(el)
    tblPr.append(b)


def table_indent(t, pt):
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), str(int(pt * 20))); ind.set(qn("w:type"), "dxa")
    t._tbl.tblPr.append(ind)


def fixed_layout(t, widths_pt):
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_pt): gc.set(qn("w:w"), str(int(w * 20)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_pt): c.width = Pt(w)


def cell(c, lines, size=10, align=None, valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER, indent=None, line=None):
    if isinstance(lines, str): lines = lines.split("\n")
    p = c.paragraphs[0]
    for i, text in enumerate(lines):
        if i: p = c.add_paragraph()
        set_font(p.add_run(text), size=size)
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        if line: pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(line)
        else: pf.line_spacing = 1.0
        if align is not None: p.alignment = align
        if indent is not None: pf.left_indent = Pt(indent)
    c.vertical_alignment = valign
    return c


def row_height(row, pt):
    row.height, row.height_rule = Pt(pt), WD_ROW_HEIGHT_RULE.AT_LEAST


def footer_logo(d):
    s = d.sections[0]
    s.footer_distance = Cm(1.75)   # 큐브티 v1 코디 양식과 같은 값 — 로고가 원본처럼 오른쪽 아래(약 761~780pt)에 온다
    fp = s.footer.paragraphs[0]
    fp.alignment = R
    fp.paragraph_format.space_after = Pt(0)
    fp.add_run().add_picture(LOGO, width=Emu(719500))   # v1 바닥글 로고와 같은 크기(2.0cm)


# ───────────────────────── 1) 주말 학습실장 근로계약서 ─────────────────────────
DAYS = ["월요일", "화요일", "수요일", "목요일", "금요일", "토요일", "일요일"]   # 일요일 칸은 본부 요청(10/23)으로 원본 표 끝에 더함
DAY_ROWS = [("근로시간", "근로시간"), ("시업", "시업시각"), ("종업", "종업시각"), ("휴게 시간", "휴게시간")]   # (표 글자, 칸 이름 끝말)


def t_weekend():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin, s.right_margin = Cm(1.27), Cm(1.13)
    s.top_margin, s.bottom_margin = Cm(1.27), Cm(2.5)
    st = d.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(10)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    footer_logo(d)

    # 제목·머리말
    para(d, [("근 로 계 약 서", True)], align=C, size=16, bold=True, line=30)
    para(d, ['넥스큐브코퍼레이션 주식회사(이하 "갑" 이라 함)와 학습실장 ', ("{직원명}", True), ' (이하 "을" 이라 함)은(는) 아래와 같이 |BR|근로계약을 체결한다.'],
         before=35.3, line=24.6)
    para(d, "― 아 래 ―", align=C, before=1.5)
    head(d, "제 1조 【근로계약기간】", before=0)
    para(d, ['"을"의 근로계약 기간은 ', ("{계약시작일}부터 {계약종료일}", True), "로 한다."])
    head(d, "제 2조 【성실의무】")
    para(d, '"을"은 근로계약의 체결에 따라 "갑"에게 성실하게 근로를 제공해야 하며, "갑"의 제반 복무규정을 준수하고 "갑"의 |BR|지시에 따라야 한다.', line=24.6)
    head(d, "제 3조 【직무 및 수행장소】", before=GAP + 1.8)
    para(d, ['가. "을"의 기본직무는 ', ("{직무}", True), "로 한다."], indent=IND)
    para(d, ['나. "을"의 직무수행 장소는 ', ("{지점}점 에듀플렉스로", True), " 한다."], indent=IND)
    para(d, '다. "을"의 직무는 "갑"의 사정상 변경될 수 있으며 이에 대한 항변은 회사 내규가 정한 절차에 따른다.', indent=IND)   # 원본은 다.만 더 들어가 있었다 — 가.·나.와 맞춤(10/23)
    head(d, "제 4조 【근무시간 및 휴게시간】")
    para(d, '가. "을"의 근로일 및 근로일별 근로시간은 다음과 같다.', indent=IND)
    para(d, "", line=21)   # 원본: 가. 줄과 표 사이 빈 줄

    t = d.add_table(rows=5, cols=len(DAYS) + 1)
    borders(t, sz=8); table_indent(t, 47.7); fixed_layout(t, [61.5] + [56.0] * len(DAYS))   # 일요일 칸이 들어가 요일 칸을 61.5→56pt 로
    for j, day in enumerate(DAYS):
        cell(t.cell(0, j + 1), day, align=C)
    for i, (label, suffix) in enumerate(DAY_ROWS):
        last = i == len(DAY_ROWS) - 1
        va = WD_CELL_VERTICAL_ALIGNMENT.TOP if last else WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell(t.cell(i + 1, 0), label, align=L, indent=6, valign=va)
        for j, day in enumerate(DAYS):
            cell(t.cell(i + 1, j + 1), "{" + day + suffix + "}", size=8, align=C, valign=va)
    for i, row in enumerate(t.rows): row_height(row, 46.5 if i == 4 else 29.8)
    for c in t.rows[4].cells: c.paragraphs[0].paragraph_format.space_before = Pt(6)   # 휴게 시간 줄은 위에 붙어 있다(원본)

    para(d, "나. 근로시간과 근무일은 “갑”과 “을”이 사전에 합의하여 사전 고지된 스케줄을 따르며,", indent=IND, before=GAP - 2)
    para(d, "사전에 합의된 내용에 따라 달라질 수 있다.", indent=43, line=24.6)
    para(d, '다. "을"의 휴게시간은 위 표에 따른다.', indent=IND, line=25.8)
    # 원본 1쪽은 여기서 끝난다(2쪽은 「라. 기타 사항 :」부터)
    p = para(d, "라. 기타 사항 :", indent=IND)
    p.paragraph_format.page_break_before = True
    para(d, ["1. 본 계약의 모든 법정 수당 산정의 기준이 되는 통상시급은 ", ("{통상시급}", True), "원으로 정한다."], indent=43)
    head(d, "제 5조 【휴일 및 연차휴가】", before=GAP + 3.3)
    para(d, '가. "을"의 휴일은 주당 1일의 주휴일(일요일)과 법정공휴일로 정한다.', indent=IND)
    para(d, "나. “을”의 휴가는 근로기준법에 따른 연차휴가가 있으며, 그 세부사항은 회사 내규에 따른다.", indent=IND)
    head(d, "제 6조 【임금】")
    para(d, '가. "을"의 근무기간 중 임금은 월급제로 한다.', indent=IND)
    para(d, "나. 월 급여 구성항목은 다음과 같다.", indent=IND, before=GAP - 0.6)

    t2 = d.add_table(rows=3, cols=4)
    borders(t2, sz=8); table_indent(t2, 27.9); fixed_layout(t2, [74.4, 91.8, 134.4, 177.6])
    hdr = t2.cell(0, 0).merge(t2.cell(0, 1))
    cell(hdr, "항 목", size=9, align=C); cell(t2.cell(0, 2), "지 급 액", size=9, align=C); cell(t2.cell(0, 3), "비 고", size=9, align=C)
    TOP = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell(t2.cell(1, 0), "월 급여", size=9, align=C, valign=TOP, line=16.3)
    cell(t2.cell(1, 1), ["① 기본급", "② 주말수당", "③ 주휴수당"], size=9, indent=5.5, valign=TOP, line=16.3)
    cell(t2.cell(1, 2), ["{기본급금액}", "{주말수당금액}", "{주휴수당금액}"], size=9, align=C, valign=TOP, line=16.3)
    cell(t2.cell(1, 3), "약 {월근로시간} H", size=9, align=C, valign=TOP, line=16.3)
    tot = t2.cell(2, 0).merge(t2.cell(2, 1))
    cell(tot, "④ 월 급여합계", size=9, indent=5.5)
    cell(t2.cell(2, 2), "{주말급여합계}", size=9, align=C)   # 서버 파생 칸 = ①+②+③(lib/contract-fields.ts)
    cell(t2.cell(2, 3), "", size=9)
    for c in (t2.cell(1, 0), t2.cell(1, 1), t2.cell(1, 2), t2.cell(1, 3)): c.paragraphs[0].paragraph_format.space_before = Pt(3)
    row_height(t2.rows[0], 22.6); row_height(t2.rows[1], 59.4); row_height(t2.rows[2], 26.4)

    para(d, ["다. 산정기간 및 지급방법 : 임금은 ", ("매월 1일부터 말일까지 ", True), "기간에 대하여 ", ("매월 말일", True), "에 “을”이 지정한 |BR|거래 은행계좌로 입금한다."],
         indent=43, first=-22.2, before=GAP + 3, line=24.6)
    head(d, "제 7조 【근태사항】", before=GAP + 1.8)
    para(d, "지각, 조퇴, 결근과 징계사항은 회사내규에 따른다.")
    head(d, "제 8조 【고용 해지】")
    para(d, '"갑"은 아래의 경우 "을"의 고용을 해지할 수 있다.')
    for it in ["가. 수습기간 중 직원으로서 부적격하다고 판단이 될 때", '나. "을"의 업무수행 능력이 부적합하다고 판단될 때', '다. "을"이 업무상 "갑"에게 손해를 끼쳤을 때',
               '라. "갑"의 제규정을 위반하여 고용관계를 유지할 수 없을 때', '마. "갑"이 경영상의 이유로 고용관계를 유지할 수 없을 때']:
        para(d, it, indent=IND)
    head(d, "제 9조 【기밀유지】")
    # 원본은 가.와 나.가 한 문단으로 붙어 있었다(「…책임을 진다.  나. "갑"과…」) — 평일판처럼 두 문단으로, 둘째 줄은 글 첫머리에 맞춤(10/23)
    para(d, '가. "을"은 근무 중 취득한 "갑"의 경영상의 지식과 정보 등 일체의 기밀을 재직 중에는 물론 퇴직 후에도 타인에게 누설하여서는 안 되며, 상기 사항의 위반으로 인한 민∙형사상의 책임을 진다.',
         indent=IND + HANG, first=-HANG, line=24.6)
    para(d, '나. "갑"과 "을"은 급여액에 대하여 타인에게 누설하지 않으며 이를 인지할 수 있는 어떠한 행동도 하지 않을 것을 확약한다.',
         indent=IND + HANG, first=-HANG, line=24.6)
    head(d, "제 10조 【준용】", before=GAP + 1.8)
    para(d, '본 계약서에 명시되지 않은 사항은 "갑"의 제규정, 노동법규 및 일반관행에 따른다.')   # 원본 「갑의」 — 문서 안 다른 곳처럼 따옴표(10/23)
    head(d, "【부칙】")
    para(d, "가. 본 계약서 이전에 작성된 고용계약서는 이 고용계약서가 체결된 즉시 효력을 잃는다.", indent=IND)
    # 날짜·서명란(굵게, 오른쪽)
    para(d, [("      {작성일}", True)], align=R, bold=True, before=69)
    para(d, "갑. 넥스큐브코퍼레이션 주식회사 고 승 재   {대표서명} (인)", align=R, bold=True, before=GAP)
    para(d, [f"{WEEKEND_BRANCH_TEXT} 에듀플렉스 원장 ", ("{원장명}", True), "   《원장서명》 (인)"], align=R, bold=True, before=GAP)
    para(d, ["을. ", ("{직원명}", True), "   {근로자서명} (인)"], align=R, bold=True, before=GAP)
    return d, "주말 학습실장 근로계약서.docx"


# ───────────────────────── 2) 코디 근로계약서(계약직 일반) 오전·오후 ─────────────────────────
V1_TIME = "기준  {출근시각} ～ {퇴근시각} 로 한다."


def codi_variant(pm):
    d = Document(CODI_V1)
    hit_time = hit_rest = 0
    for p in d.paragraphs:
        if V1_TIME in p.text:
            runs = [r for r in p.runs if V1_TIME in r.text]
            assert len(runs) == 1, "v1 제4조 나. 문장이 한 run 에 있지 않다 — 양식이 바뀌었는지 확인"
            if pm:
                a, b = PM_FIELDS
                runs[0].text = runs[0].text.replace(V1_TIME, f"기준  오후 {{{a}}} ～ 오후 {{{b}}} 로 한다.")   # 「오후 2:00 ～ 오후 6:00」
            hit_time += 1
        if p.text.startswith('다.  "을"의 휴게시간은'):
            # v1 만 들여쓰기가 빠져 있다 — 가.·나.와 같은 첫 줄 들여쓰기(2글자)
            pPr = p._p.get_or_add_pPr()
            for old in pPr.findall(qn("w:ind")): pPr.remove(old)
            ind = OxmlElement("w:ind"); ind.set(qn("w:firstLineChars"), "200"); ind.set(qn("w:firstLine"), "400")
            ps = pPr.find(qn("w:pStyle"))
            (ps.addnext(ind) if ps is not None else pPr.insert(0, ind))
            hit_rest += 1
    assert hit_time == 1 and hit_rest == 1, (hit_time, hit_rest)
    name = "코디 근로계약서(계약직 일반·오후).docx" if pm else "코디 근로계약서(계약직 일반·오전).docx"
    return d, name


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    d, name = t_weekend(); d.save(os.path.join(OUT, name)); print("saved", name)
    if os.path.exists(CODI_V1):
        for pm in (True, False):
            d, name = codi_variant(pm); d.save(os.path.join(OUT, name)); print("saved", name)
    else:
        print("코디 v1 양식을 못 찾아 코디 두 판은 건너뜀:", CODI_V1)
