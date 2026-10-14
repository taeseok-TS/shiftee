# 큐브티 추가 템플릿 8종(2026-10-13, 직영 HR팀 요청 — 모두싸인 → 큐브티) — PDF 원본을 워드 양식으로 다시 만든다.
#
# 원본은 PDF 라 치환 엔진(docxtemplater)이 못 읽는다. 문구·조항·표 숫자는 PDF 그대로 옮기고 입력 자리는 {칸} 으로 둔다.
#  · 자동 칸: {직원명} {지점} {지점명} {직급} {생년월일} {주소} {작성일} {원장명} {지점주소} {계약시작일} {계약종료일}
#  · 서명 마커: {근로자서명} 《원장서명》 《본부서명》 {대표서명}(직인)
#  · 서명 때 고르는 동의: {동의_…}  (consentBox → "(☒ 동의함  □ 동의하지 않음)")
# 실행: python build_tpl8.py <출력 폴더>
# 2026-10-14 박정인 님 답(봇 DM) 반영: 비밀유지서약서·교육 동의서의 주민등록번호 → {생년월일}, 개인정보 동의서의 고유식별정보(주민등록번호·운전면허·여권·외국인등록번호) 표·동의 칸 삭제.
#   채용 단계 서류라 고유식별정보를 받지 않는다. 1항 문구의 「고유식별정보」도 함께 뺐다.
import sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
FONT = "맑은 고딕"
C = WD_ALIGN_PARAGRAPH.CENTER
R = WD_ALIGN_PARAGRAPH.RIGHT
J = WD_ALIGN_PARAGRAPH.JUSTIFY


def new_doc(compact=False):
    # compact: 원본이 한 쪽인 양식(개인정보 동의서·퇴직금 동의서)은 글자·줄 간격을 줄여 한 쪽에 맞춘다
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8 if compact else 2.0)
    s.top_margin, s.bottom_margin = Cm(1.5 if compact else 2.0), Cm(1.2 if compact else 1.8)
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(9.5 if compact else 10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(3 if compact else 4)
    st.paragraph_format.line_spacing = 1.1 if compact else 1.25
    return d


def run(p, text, bold=False, size=None, color=None, underline=False):
    r = p.add_run(text)
    r.bold = bold
    r.underline = underline
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return r


def para(d, text="", bold=False, size=None, align=None, after=None, before=None, indent=None):
    p = d.add_paragraph()
    if text: run(p, text, bold=bold, size=size)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if indent is not None: p.paragraph_format.left_indent = Cm(indent)
    return p


def title(d, text, size=18, after=14):
    return para(d, text, bold=True, size=size, align=C, after=after)


def set_cell(cell, text, bold=False, size=None, align=None, shade=None, valign=True):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for line_i, line in enumerate(text.split("\n")):
        if line_i: p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        run(p, line, bold=bold, size=size)
        if align is not None: p.alignment = align
    if shade:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), shade)
        tcPr.append(shd)
    if valign: cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return cell


def table(d, rows, widths=None, header=False, size=10, align_center_cols=(), shade_header="EDEDED", borders=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if borders: t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            is_h = header and i == 0
            set_cell(t.cell(i, j), str(val), bold=is_h, size=size, align=C if (is_h or j in align_center_cols) else None, shade=shade_header if is_h else None)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    return t


def no_border(t):
    tbl = t._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "nil"); b.append(el)
    tblPr.append(b)


def sig_line(d, label, name_tag, marker, suffix="(서명)", align=None, before=6):
    p = para(d, before=before, after=6, align=align)
    run(p, f"{label}{name_tag}   {marker} {suffix}")
    return p


# ───────────────────────── 1) 출근일 변경 신청서(루트코스) ─────────────────────────
def t_start_change():
    d = new_doc()
    title(d, "출근일 변경 신청서", after=4)
    para(d, "※ 지점 및 사용자(원장 등) 사유/요청으로 인한 출근일 변경시 사용불가", size=9, align=C, after=12)
    para(d, "1. 대상자(근로자) 인적사항", bold=True, before=6)
    for line in ["1) 성명 : {직원명}", "2) 소속지점/직급 : {지점} / {직급}", "3) 생년월일 : {생년월일}", "4) 교육 과정명 : 매니저 루트코스", "5) 교육 수료일 : {교육수료일}"]:
        para(d, line, indent=0.5, after=2)
    para(d, "2. 출근일 변경 내용", bold=True, before=10)
    for line in ["1) 당초 출근 예정일 : {당초출근예정일}", "2) 변경 출근일 : {변경출근일}"]:
        para(d, line, indent=0.5, after=2)
    para(d, "3. 출근일 변경 신청 사유", bold=True, before=10)
    para(d, "본인은 개인 사정(이사, 건강, 가족 돌봄 등)으로 인해 당초 예정된 출근일에 근로 제공이 어려운 상황입니다. 이에 따라 출근일을 위와 같이 변경하여 줄 것을 본인의 의사에 따라 자발적으로 신청합니다.", align=J, indent=0.5)
    para(d, "4. 회사 안내 및 근로자 선택 확인", bold=True, before=10)
    para(d, "회사는 당초 예정된 출근일에 정상적으로 근로 제공이 가능함을 안내하였으며, 본인은 개인 사정에 따라 출근일 변경을 선택하였음을 확인합니다.", align=J, indent=0.5)
    para(d, "5. 임금 및 근로조건 관련 확인", bold=True, before=10)
    para(d, "변경된 출근일 이전 기간 동안에는 근로 제공이 이루어지지 않으므로 해당 기간에 대한 임금은 발생하지 않음을 충분히 이해합니다. 또한 해당 기간은 실제 근로 제공이 없는 기간임을 확인합니다.", align=J, indent=0.5)
    para(d, "본 신청은 본인의 자유로운 의사에 따라 이루어진 것이며 회사로부터 강요 또는 불이익 우려 없이 작성되었음을 확인합니다.", align=J, indent=0.5, after=18)
    d.add_page_break()   # 원본처럼 날짜·서명란은 2쪽에 모아 둔다(서명 세 곳이 한 쪽에)
    para(d, "{작성일}", align=C, before=12, after=14)
    sig_line(d, "신청자(근로자) 성명: ", "{직원명}", "{근로자서명}")
    sig_line(d, "지점 원장: ", "{원장명}", "《원장서명》")
    sig_line(d, "직영사업본부 확인: ", "", "《본부서명》", suffix="(인)")
    para(d, "넥스큐브코퍼레이션  귀중", bold=True, before=14)
    return d, "출근일 변경 신청서(루트코스).docx"


# ───────────────────────── 2) 육아휴직원(매니저·교실장) ─────────────────────────
def t_parental_leave():
    d = new_doc()
    # 머리: 왼쪽 회사·지점, 오른쪽 결재표(원장|본부)
    head = d.add_table(rows=1, cols=2); no_border(head)
    left, right = head.cell(0, 0), head.cell(0, 1)
    left.width, right.width = Cm(10.5), Cm(6.5)
    set_cell(left, "넥스큐브코퍼레이션㈜\n에듀플렉스 {지점명} 직영점", size=10, valign=False)
    appr = right.add_table(rows=2, cols=3); appr.style = "Table Grid"; appr.alignment = WD_TABLE_ALIGNMENT.RIGHT
    set_cell(appr.cell(0, 0), "결\n재", bold=True, size=9, align=C, shade="EDEDED"); appr.cell(0, 0).merge(appr.cell(1, 0))
    set_cell(appr.cell(0, 1), "원장", bold=True, size=9, align=C, shade="EDEDED"); set_cell(appr.cell(0, 2), "본부", bold=True, size=9, align=C, shade="EDEDED")
    set_cell(appr.cell(1, 1), "《원장서명》", size=9, align=C); set_cell(appr.cell(1, 2), "《본부서명》", size=9, align=C)
    for r_ in appr.rows: r_.height = Cm(1.2)
    for row in appr.rows:
        row.cells[0].width = Cm(1.0); row.cells[1].width = Cm(2.6); row.cells[2].width = Cm(2.6)
    right.paragraphs[0].text = ""
    title(d, "육 아 휴 직 원", after=8)
    para(d, "상기 본인은 다음과 같이 육아 휴직하고자 하오니 허가하여 주시기 바랍니다.", align=C, after=10)
    t = d.add_table(rows=8, cols=6); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    S = "EDEDED"
    # 인적사항 (세로 머리글 병합)
    set_cell(t.cell(0, 0), "인  적\n사  항", bold=True, align=C, shade=S); t.cell(0, 0).merge(t.cell(4, 0))
    set_cell(t.cell(0, 1), "성  명", bold=True, align=C, shade=S); set_cell(t.cell(0, 2), "{직원명}", align=C)
    set_cell(t.cell(0, 3), "소  속", bold=True, align=C, shade=S); set_cell(t.cell(0, 4), "{지점}", align=C)
    set_cell(t.cell(0, 5), "직  급\n{직급}", align=C)
    set_cell(t.cell(1, 1), "주소", bold=True, align=C, shade=S); set_cell(t.cell(1, 2), "{주소}"); t.cell(1, 2).merge(t.cell(1, 5))
    set_cell(t.cell(2, 1), "연락처", bold=True, align=C, shade=S); set_cell(t.cell(2, 2), "{본인연락처}"); t.cell(2, 2).merge(t.cell(2, 3))
    set_cell(t.cell(2, 4), "생년월일", bold=True, align=C, shade=S); set_cell(t.cell(2, 5), "{생년월일}", align=C)
    set_cell(t.cell(3, 1), "부재 중\n연락처\n(본인 외)", bold=True, align=C, shade=S); set_cell(t.cell(3, 2), "{긴급연락처}   (관계: {관계})"); t.cell(3, 2).merge(t.cell(3, 5))
    set_cell(t.cell(4, 1), "", shade=S); t.cell(4, 1).merge(t.cell(4, 5))
    # 내용
    set_cell(t.cell(5, 0), "내  용", bold=True, align=C, shade=S); t.cell(5, 0).merge(t.cell(7, 0))
    set_cell(t.cell(5, 1), "휴직 사유", bold=True, align=C, shade=S); set_cell(t.cell(5, 2), "자녀 양육을 위한 육아휴직 신청"); t.cell(5, 2).merge(t.cell(5, 5))
    set_cell(t.cell(6, 1), "휴직시작일", bold=True, align=C, shade=S); set_cell(t.cell(6, 2), "{휴직시작일}", align=C); t.cell(6, 2).merge(t.cell(6, 3))
    set_cell(t.cell(6, 4), "휴직종료일", bold=True, align=C, shade=S); set_cell(t.cell(6, 5), "{휴직종료일}", align=C)
    set_cell(t.cell(7, 1), "복직예정일", bold=True, align=C, shade=S); set_cell(t.cell(7, 2), "{복직예정일}", align=C); t.cell(7, 2).merge(t.cell(7, 5))
    for row in t.rows:
        row.height = Cm(1.0)
        ws = [1.8, 2.6, 3.4, 2.6, 3.0, 3.6]
        for j, w in enumerate(ws): row.cells[j].width = Cm(w)
    para(d, "", after=14)
    para(d, "신청일 : {작성일}", align=R, after=6)
    p = para(d, align=R, after=6); run(p, "신청인 : {직원명}   {근로자서명} (인)")
    return d, "육아휴직원(매니저·교실장).docx"


# ───────────────────────── 3) [채용] 개인정보 수집·이용 동의서 ─────────────────────────
def t_privacy_hire():
    d = new_doc(compact=True)
    title(d, "개인정보의 수집 ∙ 이용에 관한 동의서", size=16, after=10)
    para(d, "1. {직원명} 은(는) 에듀플렉스 {지점명} 지점 재직근로자로서 인사관리상 개인정보의 수집 ∙ 이용이 필요하다는 것을 이해하고 있고 다음과 같은 개인정보 ∙ 민감정보를 수집 ∙ 이용하는 것에 동의합니다.", align=J)
    table(d, [["개인정보항목", "수집∙ 이용 목적", "보유기간"],
              ["가. 성명\n나. 주소, 이메일, 연락처\n다. 학력, 근무경력, 자격증\n라. 기타 근무와 관련된 개인정보", "가. 채용 및 승진 등 인사관리\n나. 세법, 노동관계 법령 등에서 부과하는 의무이행", "재직기간 동안 보유하고, 기타 개별 법령에서 보유기간을 정하고 있는 경우 그에 따름"]],
          widths=[6.0, 5.5, 5.5], header=True, size=9.5)
    para(d, "개인정보의 수집 ∙ 이용에 {동의_개인정보 수집이용}", align=R, before=4, after=10)
    table(d, [["민감정보의 항목", "수집 ∙ 이용 목적", "보유기간"],
              ["가. 신체장애\n나. 병력\n다. 범죄정보", "가. 채용 및 승진 등 인사관리\n나. 세법, 노동관계법령 등에서 부과하는 의무이행\n다. 정부지원금 신청", "재직기간 동안 보유하고, 기타 개별 법령에서 보유기간을 정하고 있는 경우 그에 따름"]],
          widths=[6.0, 5.5, 5.5], header=True, size=9.5)
    para(d, "민감정보의 수집 ∙ 이용에 {동의_민감정보 수집이용}", align=R, before=4, after=10)
    para(d, "2. {직원명} 은(는) 에듀플렉스 {지점명} 지점이 취득한 개인정보를 재직기간 동안 내부적으로 채용 ∙ 승진 등 인사관리에 이용하고, 외부적으로 법령에 따라 관계기관에 제공하는 것에 동의합니다.", align=J)
    para(d, "개인정보의 수집 ∙ 이용에 {동의_내부이용 외부제공}", align=R, before=2, after=10)
    para(d, "3. 에듀플렉스 {지점명} 지점은 취득한 개인정보를 수집한 목적에 필요한 범위에서 적합하게 처리하고 그 목적 외의 용도로 사용하지 않으며 개인 정보를 제공한 계약당사자는 언제나 자신이 입력한 개인정보를 열람 ∙ 수정 및 정보제공에 대한 철회를 할 수 있습니다.", align=J)
    para(d, "4. 본인은 1~3항에 따라 수집되는 개인정보의 항목과 개인정보의 수집 ∙ 이용에 대한 거부를 할 수 있는 권리가 있다는 사실을 충분히 설명 받고 숙지하였으며, 미동의시 적법하게 시행되는 회사내부규정 및 법령에 따라 발생하는 불이익에 대한 책임은 본인에게 있음을 확인합니다.", align=J, after=16)
    para(d, "{작성일}", align=C, after=10)
    p = para(d, align=R, after=8); run(p, "동의자 성명: {직원명}   {근로자서명} (인)")
    para(d, "에듀플렉스 {지점명} 지점   원장 {원장명}   귀중", bold=True, align=C, before=8)
    return d, "[채용] 개인정보 수집이용 동의서.docx"


# ───────────────────────── 4) [채용] EDUPLEX 비밀유지서약서 ─────────────────────────
def t_nda_hire():
    d = new_doc()
    title(d, "비밀유지서약서", size=18, after=14)
    para(d, "성 명 : {직원명}            (생년월일 : {생년월일})", after=4)
    para(d, "주 소 : {주소}", after=12)
    para(d, "상기 본인은 고용계약서 제12조 제3항에 의거, {지점명} 지점과 고용 계약기간 중 얻게 된 다음 업무상 비밀을 인지하고 준수하며, 계약이 종료된 이후에도 2년간 동종업계에서의 창업 및 이∙전직으로 업무상의 비밀 누설을 하지 않을 것임을 서약합니다. 이를 위반할 경우, 관련 법규 및 고용계약서에 의거 {지점명} 지점 원장과 에듀플렉스 본사는 상기 본인에게 손해배상을 청구할 수 있음을 확인합니다. 이에 아래의 사항에 대하여 서약합니다.", align=J, after=12)
    para(d, "-다 음-", bold=True, align=C, after=10)
    para(d, "1. 고용계약 기간 중, “에듀플렉스”로부터 공급받은 매뉴얼, 매니지먼트 교재, 테스트, EMS 및 ECHOS 등의 경영시스템, 각종 교육 자료 및 세미나 자료 등에 관련된 사항에 대해서 비밀유지의무를 준수할 것임을 서약합니다.", align=J, after=8)
    para(d, "2. 고용계약 기간 중, 습득한 직원 및 회원(학생/학부모)의 개인정보와 관련된 사항에 대한 비밀유지의무를 준수할 것임을 서약합니다.", align=J, after=24)
    para(d, "{작성일}", align=C, after=14)
    p = para(d, align=R, after=14); run(p, "서약인 : {직원명}   {근로자서명} (인)")
    para(d, "에듀플렉스 ( {지점명} ) 지점 원장 ( {원장명} ) 귀중", bold=True, align=C, after=4)
    para(d, "넥스큐브코퍼레이션 주식회사 대표이사 고승재 귀중", bold=True, align=C)
    return d, "[채용] 비밀유지서약서.docx"


# ───────────────────────── 5) 부관리자 성과급 계약서 ─────────────────────────
INCENTIVE = [
    (21, "7.0", "30,000", "630,000"), (22, "7.3", "30,000", "660,000"), (23, "7.7", "30,000", "690,000"), (24, "8.0", "30,000", "720,000"), (25, "8.3", "30,000", "750,000"),
    (26, "8.7", "30,000", "780,000"), (27, "9.0", "30,000", "810,000"), (28, "9.3", "30,000", "840,000"), (29, "9.7", "30,000", "870,000"), (30, "10.0", "30,000", "900,000"),
    (31, "10.3", "31,000", "961,000"), (32, "10.7", "31,000", "992,000"), (33, "11.0", "31,000", "1,023,000"), (34, "11.3", "31,000", "1,054,000"), (35, "11.7", "31,000", "1,085,000"),
    (36, "12.0", "31,000", "1,116,000"), (37, "12.3", "31,000", "1,147,000"), (38, "12.7", "31,000", "1,178,000"), (39, "13.0", "31,000", "1,209,000"), (40, "13.3", "31,000", "1,240,000"),
    (41, "13.7", "32,000", "1,312,000"), (42, "14.0", "32,000", "1,344,000"), (43, "14.3", "32,000", "1,376,000"), (44, "14.7", "32,000", "1,408,000"), (45, "15.0", "32,000", "1,440,000"),
    (46, "15.3", "32,000", "1,472,000"), (47, "15.7", "32,000", "1,504,000"), (48, "16.0", "32,000", "1,536,000"), (49, "16.3", "32,000", "1,568,000"), (50, "16.7", "32,000", "1,600,000"),
    (51, "17.0", "35,000", "1,785,000"), (52, "17.3", "35,000", "1,820,000"), (53, "17.7", "35,000", "1,855,000"), (54, "18.0", "35,000", "1,890,000"), (55, "18.3", "35,000", "1,925,000"),
    (56, "18.7", "35,000", "1,960,000"), (57, "19.0", "35,000", "1,995,000"), (58, "19.3", "35,000", "2,030,000"), (59, "19.7", "35,000", "2,065,000"), (60, "20.0", "35,000", "2,100,000"),
    (61, "20.3", "35,000", "2,135,000"), (62, "20.7", "35,000", "2,170,000"), (63, "21.0", "35,000", "2,205,000"), (64, "21.3", "35,000", "2,240,000"), (65, "21.7", "35,000", "2,275,000"),
    (66, "22.0", "35,000", "2,310,000"), (67, "22.3", "35,000", "2,345,000"), (68, "22.7", "35,000", "2,380,000"), (69, "23.0", "35,000", "2,415,000"), (70, "23.3", "35,000", "2,450,000"),
    (71, "23.7", "40,000", "2,840,000"), (72, "24.0", "40,000", "2,880,000"), (73, "24.3", "40,000", "2,920,000"), (74, "24.7", "40,000", "2,960,000"), (75, "25.0", "40,000", "3,000,000"),
    (76, "25.3", "40,000", "3,040,000"), (77, "25.7", "40,000", "3,080,000"), (78, "26.0", "40,000", "3,120,000"), (79, "26.3", "40,000", "3,160,000"), (80, "26.7", "40,000", "3,200,000"),
    (81, "27.0", "40,000", "3,240,000"), (82, "27.3", "40,000", "3,280,000"), (83, "27.7", "40,000", "3,320,000"), (84, "28.0", "40,000", "3,360,000"), (85, "28.3", "40,000", "3,400,000"),
    (86, "28.7", "40,000", "3,440,000"), (87, "29.0", "40,000", "3,480,000"), (88, "29.3", "40,000", "3,520,000"), (89, "29.7", "40,000", "3,560,000"), (90, "30.0", "45,000", "4,050,000"),
    (91, "30.3", "45,000", "4,095,000"), (92, "30.7", "45,000", "4,140,000"), (93, "31.0", "45,000", "4,185,000"), (94, "31.3", "45,000", "4,230,000"), (95, "31.7", "45,000", "4,275,000"),
    (96, "32.0", "45,000", "4,320,000"), (97, "32.3", "45,000", "4,365,000"), (98, "32.7", "45,000", "4,410,000"), (99, "33.0", "45,000", "4,455,000"), (100, "33.3", "50,000", "5,000,000"),
]


def t_incentive():
    d = new_doc()
    title(d, "성과급 계약서", after=12)
    para(d, "넥스큐브코퍼레이션 주식회사(이하 “갑”이라 함)와 직원 {직원명} 은/는(이하 “을”이라 함) 아래와 같이 성과급계약을 체결하고 상호 성실히 이행 준수할 것을 서약합니다.", align=J, after=10)
    para(d, "제1조 (계약기간)", bold=True, before=6, after=2)
    para(d, "성과급계약 기간은 {계약시작일}부터 {계약종료일} 까지로 한다.", after=8)
    para(d, "제2조 (성과급기준)", bold=True, before=6, after=2)
    para(d, "① 성과급의 기준은 “을”이 근무하고 있는 해당지점에서, “을”이 등록상담 진행 후 신규등록 한 수와 그에 따른 기준금액으로 한다.", align=J, after=2)
    para(d, "② 성과급 산정 기준은 다음과 같다.", after=6)
    rows = [["3개월 등록수", "1개월 평균", "기준금액", "인센티브 금액"]] + [[str(a), b, c_, e] for a, b, c_, e in INCENTIVE]
    t = table(d, rows, widths=[4.0, 4.0, 4.0, 4.5], header=True, size=9.5, align_center_cols=(0, 1, 2, 3))
    for row in t.rows: row.height = Cm(0.55)
    para(d, "", after=6)
    para(d, "제3조 (지급방법)", bold=True, before=8, after=2)
    para(d, "성과급은 결산이 끝난 날의 익월 임금에 포함하여 지급하며, 지급일 기준으로 재직중인 자에게만 지급하는 것을 원칙으로 한다.", align=J, after=8)
    para(d, "제4조 (계약갱신)", bold=True, before=6, after=2)
    para(d, "① 별도의 고지가 없는 경우 이전에 적용한 성과급계약 조건과 동일하게 계약이 갱신된 것으로 본다.", align=J, after=2)
    para(d, "② 승진, 휴직(징계) 및 보직변경 등 “갑”의 규정에 있는 사유로 연봉조정이 불가피한 경우는 별도로 정한 바에 의하여 조정한다.", align=J, after=8)
    para(d, "제5조 (계약해지)", bold=True, before=6, after=2)
    para(d, "고용계약서에 표시된 사유로 고용을 해지할 시 본 계약도 자연히 해지한다.", after=8)
    para(d, "제6조 (기밀유지)", bold=True, before=6, after=2)
    para(d, "“갑”과 “을”은 지점 성과급 계약 금액에 대하여 타인에게 누설하지 않으며 이를 인지할 수 있는 어떠한 행동도 하지 않을 것을 확약한다.", align=J, after=8)
    para(d, "제7조 (준용)", bold=True, before=6, after=2)
    para(d, "본 계약서에 명기되지 않은 사항은 “갑”의 사규 및 제규정 등에 따른다.", after=24)
    para(d, "{계약일}", align=C, after=20)
    p = para(d, after=14); run(p, "갑. 넥스큐브코퍼레이션 대표이사  고 승 재   {대표서명} (인)")
    p = para(d, after=6); run(p, "을. 에듀플렉스 {지점명}점 {직급}  {직원명}   {근로자서명} (인)")
    return d, "부관리자 성과급 계약서.docx"


# ───────────────────────── 6) [채용] 채용과정 교육 동의서 ─────────────────────────
def t_training_consent():
    d = new_doc()
    title(d, "동 의 서", after=14)
    para(d, "성 명 : {직원명}            (생년월일 : {생년월일})", after=4)
    para(d, "주 소 : {주소}", after=14)
    para(d, "상기 본인은 하단의 교육 및 채용 사항에 대해 동의합니다.", after=4)
    para(d, "만약 이 동의서에 동의하지 않거나, 위반하여 귀사에 손해를 끼친 경우에는 채용확정 이후라 해도 채용이 취소될 수 있습니다.", align=J, after=12)
    items = [
        "1.  후보자 본인은 에듀플렉스의 채용단계 중 면접단계 이후 약 3주간 온라인 본사 교육이 진행되는 사항을 확인하였으며, 이에 동의합니다.",
        "2.  후보자 본인은 에듀플렉스의 약 3주간 진행되는 온라인 본사 교육의 수료자에 한하여 최종 합격 여부 및 입사여부가 결정되는 사항을 확인하였으며, 이에 동의합니다.",
        "3.  후보자 본인은 에듀플렉스의 약 3주간 진행되는 온라인 본사교육은 채용 프로세스의 일부임을 확인하였으며, 면접 단계 이후, 입사 단계 이전 진행되는 교육 절차에 동의합니다.",
        "4.  후보자는 에듀플렉스 교육과정 수료자에 한하여 지급되는 교육참가비에 대한 내용을 안내받았으며, 이에 동의합니다.",
    ]
    for it in items: para(d, it, align=J, after=8)
    para(d, "후보자 본인은 상기의 내용을 모두 인지하였으며, 이에 동의합니다.", before=10, after=6)
    para(d, "{동의_채용과정 교육}", align=C, after=20)
    para(d, "{작성일}", align=C, after=14)
    p = para(d, align=R); run(p, "동 의 인 :  {직원명}   {근로자서명} (인)")
    return d, "[채용] 채용과정 교육 동의서.docx"


# ───────────────────────── 7) 퇴직금 일반계좌 지급 요청 및 동의서 ─────────────────────────
def t_severance_account():
    d = new_doc(compact=True)
    appr = d.add_table(rows=2, cols=2); appr.style = "Table Grid"; appr.alignment = WD_TABLE_ALIGNMENT.RIGHT
    set_cell(appr.cell(0, 0), "본인", bold=True, size=9, align=C, shade="EDEDED"); set_cell(appr.cell(0, 1), "본부(부서)", bold=True, size=9, align=C, shade="EDEDED")
    set_cell(appr.cell(1, 0), "{근로자서명}\n(서명/인)", size=8, align=C); set_cell(appr.cell(1, 1), "《본부서명》\n(서명/인)", size=8, align=C)
    for row in appr.rows:
        for c_ in row.cells: c_.width = Cm(2.8)
    appr.rows[1].height = Cm(1.4)
    title(d, "퇴직금 일반계좌 지급 요청 및 동의서", size=16, after=12)
    para(d, "1. 신청인(근로자) 정보", bold=True, after=2)
    for line in ["- 성    명 : {직원명}", "- 소    속 : {지점}", "- 퇴사일자 : {퇴사일자}"]: para(d, line, indent=0.5, after=2)
    para(d, "2. 일반계좌 지급 요청 사유", bold=True, before=10, after=2)
    para(d, "본인은 「근로자퇴직급여보장법」 제9조 제2항에 의거하여 퇴직금을 개인형퇴직연금(IRP) 계좌로 수령해야 함을 고지받았으나, 본인의 불가피한 금융거래 상황으로 인하여 아래와 같이 일반계좌 지급을 요청합니다.", align=J, indent=0.5, after=4)
    para(d, "- 지급요청사유 : 금융권 대출 관련 신규계좌(IRP) 개설 원천 불가", indent=0.5, after=2)
    para(d, "- 상세사유 :", indent=0.5, after=2)
    para(d, "현재 본인의 대출실행 규정 및 ’단기간 다수계좌 개설 제한’ 제도에 따라, 해당 금융기관으로부터 신규 IRP 계좌 개설이 불가능함을 공식 확인받았습니다. 이에 따라 법정 기일 내 퇴직금 수령을 위해 부득이하게 일반계좌로의 지급을 신청합니다.", align=J, indent=0.5)
    para(d, "3. 수령 희망 계좌 정보", bold=True, before=10, after=2)
    for line in ["- 은 행 명 : {은행명}", "- 계좌번호 : {계좌번호}", "- 예 금 주 : {예금주} (본인)"]: para(d, line, indent=0.5, after=2)
    para(d, "4. 근로자 확인 및 동의사항", bold=True, before=10, after=2)
    for line in ["1. 본인은 IRP 계좌 미개설에 따른 일반계좌 지급 요청이 본인의 자발적 의사에 의한 것임을 확인합니다.",
                 "2. 일반계좌 지급 시 발생하는 퇴직소득세 원천징수 및 관련 행정절차에 동의합니다.",
                 "3. IRP 계좌 미이전에 따른 과세이연 혜택 소멸 등의 경제적 결과가 본인의 선택에 의한 것임을 인지하였으며, 이에 따른 제반 사항을 수용합니다."]:
        para(d, line, align=J, indent=0.5, after=2)
    para(d, "[참고사항]", bold=True, size=9, before=10, after=2)
    para(d, "1. 퇴직소득세 원천징수: 퇴직금을 IRP가 아닌 일반계좌로 수령할 경우, 회사가 관련 법령에 따라 해당 세액을 미리 공제하고 남은 차액만을 지급하게 됩니다.", size=9, align=J, indent=0.5, after=2)
    para(d, "2. 과세이연 혜택 소멸: IRP 계좌 수령 시에는 세금 납부가 유예(과세이연)되고 추후 연금 수령 시 세액 감면 혜택을 받을 수 있으나, 일반계좌 수령 시에는 이러한 혜택이 모두 소멸됩니다.", size=9, align=J, indent=0.5, after=16)
    para(d, "{작성일}", align=C, after=12)
    p = para(d, align=R, after=10); run(p, "신청인 : {직원명}   {근로자서명} (인/서명)")
    para(d, "넥스큐브코퍼레이션 주식회사  귀중", bold=True, align=C)
    return d, "퇴직금 일반계좌 지급 요청 및 동의서.docx"


# ───────────────────────── 8) 인사발령통지서 ─────────────────────────
def t_transfer_notice():
    d = new_doc()
    title(d, "인 사 발 령 통 지 서", after=16)
    para(d, "1. 대상자 인적사항", bold=True, after=4)
    for line in ["성 명 : {직원명}", "소 속 : 에듀플렉스 {지점명}점", "직 급 : {직급}"]: para(d, line, indent=0.7, after=3)
    para(d, "2. 발령 내용", bold=True, before=10, after=4)
    for line in ["현 근무지(전) : {지점주소}", "발령 근무지(후) : {발령근무지주소}  ({발령지점})", "발 령 일 자 : {발령일} 부"]: para(d, line, indent=0.7, after=3)
    para(d, "3. 발령 사유", bold=True, before=10, after=4)
    para(d, "{발령사유}", indent=0.7, after=3)
    para(d, "4. 기타 안내사항", bold=True, before=10, after=4)
    para(d, "본 발령에 따른 업무 인수인계는 {인수인계완료일}까지 완료하여 주시기 바랍니다.", indent=0.7, after=3)
    para(d, "본 인사발령과 관련하여 통근 거리 문제 등 고충이 있는 경우, 인사 담당자 및 원장에게 면담을 요청하시기 바랍니다.", align=J, indent=0.7, after=10)
    para(d, "위와 같이 인사를 발령함을 통지합니다.", before=6, after=24)
    para(d, "{작성일}", align=C, after=10)
    para(d, "넥스큐브코퍼레이션", bold=True, align=C, after=2)
    p = para(d, align=C, after=24); run(p, "대표이사 고승재   {대표서명} (인)", bold=True)
    para(d, "[수령 확인서]", bold=True, before=16, after=6)
    para(d, "본인은 위 인사발령 내용을 {수령확인일}에 정식으로 통지받았음을 확인합니다.", align=J, after=8)
    para(d, "확인 일자 : {수령확인일}", after=4)
    p = para(d, after=4); run(p, "성 명 : {직원명}   {근로자서명} (서명 또는 인)")
    return d, "인사발령통지서.docx"


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for fn in (t_start_change, t_parental_leave, t_privacy_hire, t_nda_hire, t_incentive, t_training_consent, t_severance_account, t_transfer_notice):
        d, name = fn()
        d.save(os.path.join(OUT, name))
        print("saved", name)
